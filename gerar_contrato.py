# -*- coding: utf-8 -*-
"""
Gerador de Instrumento Particular de Confissao de Divida + Notas Promissorias.

Uso:
    python gerar_contrato.py                -> gera o PDF de TODAS as linhas da planilha
    python gerar_contrato.py 3               -> gera o PDF apenas da linha 3 (2a linha de dados, pois a linha 1 e cabecalho)
    python gerar_contrato.py --planilha outro_arquivo.xlsx

Os PDFs sao salvos em ./saida/
"""

import sys
import re
import unicodedata
from datetime import date
from calendar import monthrange
from pathlib import Path

import openpyxl
from num2words import num2words
from fpdf import FPDF

PASTA = Path(__file__).parent
PLANILHA_PADRAO = PASTA / "clientes.xlsx"
PASTA_SAIDA = PASTA / "saida"

# ---------------------------------------------------------------------------
# Dados fixos do CREDOR (nao mudam de cliente para cliente).
# Se um dia o credor mudar, edite so aqui.
# ---------------------------------------------------------------------------
CREDOR_NOME = "Sartorato Informações Cadastrais Ltda. ME"
# O modelo grafa o credor de formas diferentes na assinatura e na nota promissória.
CREDOR_NOME_ASSINATURA = "Sartorato Informações Cadastrais LTDA. ME"
CREDOR_NOME_NOTA = "SARTORATO INFORMAÇÕES CADASTRAIS LTDA ME"
CREDOR_CNPJ = "06.301.745/0001-77"
CREDOR_ENDERECO = "Rua Benedito Américo de Oliveira, nº 325, Vila Yara, Osasco/SP, CEP 06028-080"
CREDOR_SOCIO_NOME = "Rogério Sartorato"
CREDOR_SOCIO_QUALIFICACAO = "brasileiro, solteiro, empresário"
CREDOR_SOCIO_CPF = "206.301.208-38"
CREDOR_SOCIO_RG = "23.000.034-4"
CREDOR_PIX = CREDOR_CNPJ

# Valores padrao dos encargos, usados quando a planilha nao preencher a coluna.
PADRAO_JUROS_MORA_PCT = 1
PADRAO_MULTA_MORATORIA_PCT = 2
PADRAO_MULTA_PCT = 30
PADRAO_HONORARIOS_PCT = 20
PADRAO_DESCONTO_RENDIMENTOS_PCT = 30
PADRAO_LOCAL = "Osasco"
PADRAO_FORO = "Osasco/SP"
PADRAO_TESTEMUNHA1_NOME = "Caio Vinicius Antonino Rodrigues"
PADRAO_TESTEMUNHA1_CPF = "415.510.728-74"
PADRAO_TESTEMUNHA2_NOME = "Nilza de Oliveira Sartorato"
PADRAO_TESTEMUNHA2_CPF = "256.604.838-01"

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


# ---------------------------------------------------------------------------
# Utilidades de "por extenso"
# ---------------------------------------------------------------------------

def valor_por_extenso(valor: float) -> str:
    """Ex.: 1610.00 -> 'um mil, seiscentos e dez reais'"""
    reais = int(valor)
    centavos = round((valor - reais) * 100)
    texto_reais = num2words(reais, lang="pt_BR")
    if 1000 <= reais < 2000:
        # convenção bancária: sempre explicitar "um" antes de "mil" (evita fraude/ambiguidade)
        texto_reais = "um " + texto_reais
    moeda = "real" if reais == 1 else "reais"
    texto = f"{texto_reais} {moeda}"
    if centavos:
        texto_centavos = num2words(centavos, lang="pt_BR")
        moeda_centavos = "centavo" if centavos == 1 else "centavos"
        texto += f" e {texto_centavos} {moeda_centavos}"
    return texto


def numero_por_extenso(n: int) -> str:
    return num2words(n, lang="pt_BR")


def percentual_por_extenso(pct) -> str:
    """Ex.: 30 -> '30% (trinta por cento)'"""
    inteiro = int(pct)
    if inteiro == pct:
        return f"{inteiro}% ({num2words(inteiro, lang='pt_BR')} por cento)"
    return f"{pct}% ({num2words(pct, lang='pt_BR')} por cento)"


def data_por_extenso(d: date) -> str:
    dia = num2words(d.day, lang="pt_BR").upper()
    mes = MESES[d.month - 1].upper()
    ano = num2words(d.year, lang="pt_BR").upper()
    return f"{dia} de {mes} de {ano}"


def formata_moeda(valor: float) -> str:
    s = f"{valor:,.2f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def formata_data(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def formata_data_mes_nome(d: date) -> str:
    """Ex.: 04 de Setembro de 2026"""
    return f"{d.day:02d} de {MESES[d.month - 1].capitalize()} de {d.year}"


def add_meses(d: date, n: int) -> date:
    """Soma n meses a uma data, preservando o dia (ajusta para o ultimo dia do mes se necessario)."""
    mes_total = d.month - 1 + n
    ano = d.year + mes_total // 12
    mes = mes_total % 12 + 1
    ultimo_dia = monthrange(ano, mes)[1]
    dia = min(d.day, ultimo_dia)
    return date(ano, mes, dia)


# Fonte com suporte a Unicode, necessária para o travessão "–" do modelo.
# Arial (Windows, no desenvolvimento) e Liberation Sans (container Linux) são
# metricamente compatíveis, então o layout sai igual nos dois ambientes.
FONTES_CANDIDATAS = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]


def registrar_fonte(pdf) -> str:
    """Registra a primeira fonte Unicode disponível; devolve o nome a usar.

    Sem nenhuma delas, cai para a Helvetica embutida (que não suporta travessão,
    daí a substituição em limpar_para_pdf).
    """
    for regular, negrito in FONTES_CANDIDATAS:
        if Path(regular).exists() and Path(negrito).exists():
            pdf.add_font("Documento", "", regular)
            pdf.add_font("Documento", "B", negrito)
            return "Documento"
    return "Helvetica"


def limpar_para_pdf(texto: str, fonte: str) -> str:
    """A Helvetica embutida do fpdf2 só suporta Latin-1; nesse caso troca os
    caracteres tipográficos pelos equivalentes ASCII para não quebrar."""
    if fonte != "Helvetica":
        return texto
    subs = {
        "–": "-", "—": "-",
        "‘": "'", "’": "'",
        "“": '"', "”": '"',
        "…": "...",
    }
    for k, v in subs.items():
        texto = texto.replace(k, v)
    return texto


def slugify(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^\w\s-]", "", texto).strip().lower()
    return re.sub(r"[-\s]+", "_", texto)


CONECTIVOS_NOME = {"da", "de", "do", "das", "dos", "e", "di", "du", "del", "van", "von", "y"}


def formatar_nome_proprio(nome: str) -> str:
    """'TALITA AZZI BUTTINI' -> 'Talita Azzi Buttini' (preposições em minúsculo)."""
    palavras = str(nome).strip().split()
    formatadas = []
    for i, palavra in enumerate(palavras):
        minuscula = palavra.lower()
        if i > 0 and minuscula in CONECTIVOS_NOME:
            formatadas.append(minuscula)
        else:
            formatadas.append("-".join(p[:1].upper() + p[1:] for p in minuscula.split("-")))
    return " ".join(formatadas)


# ---------------------------------------------------------------------------
# Leitura da planilha
# ---------------------------------------------------------------------------

def valor_celula(row, headers, nome, obrigatorio=True, padrao=None):
    if nome not in headers:
        if obrigatorio and padrao is None:
            raise ValueError(f"Coluna obrigatória ausente na planilha: {nome}")
        return padrao
    v = row[headers[nome]]
    if v is None or (isinstance(v, str) and v.strip() == ""):
        if obrigatorio and padrao is None:
            raise ValueError(f"Campo obrigatório vazio: {nome}")
        return padrao
    return v


def carregar_clientes(caminho_planilha: Path):
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=False))
    if not linhas:
        raise ValueError("Planilha vazia.")
    cabecalho = [str(c.value).strip() if c.value else "" for c in linhas[0]]
    headers = {nome: i for i, nome in enumerate(cabecalho) if nome}

    clientes = []
    for idx, linha_cells in enumerate(linhas[1:], start=2):
        row = [c.value for c in linha_cells]
        if all(v is None or str(v).strip() == "" for v in row):
            continue  # linha em branco, ignora

        nome_devedor = valor_celula(row, headers, "nome_devedor")

        # datas podem vir como datetime (Excel) ou string dd/mm/aaaa
        def pega_data(campo, padrao=None):
            v = valor_celula(row, headers, campo, obrigatorio=padrao is None, padrao=padrao)
            if isinstance(v, date):
                return v
            if hasattr(v, "date"):
                return v.date()
            if isinstance(v, str):
                d, m, a = v.strip().split("/")
                return date(int(a), int(m), int(d))
            return v

        cliente = {
            "linha_planilha": idx,
            "nome_devedor": nome_devedor,
            "nacionalidade": valor_celula(row, headers, "nacionalidade", padrao="brasileiro(a)"),
            "estado_civil": valor_celula(row, headers, "estado_civil", padrao="solteira"),
            "profissao": valor_celula(row, headers, "profissao"),
            "cpf_devedor": valor_celula(row, headers, "cpf_devedor"),
            "endereco_residencial": valor_celula(row, headers, "endereco_residencial"),
            "endereco_profissional": valor_celula(row, headers, "endereco_profissional", obrigatorio=False, padrao=""),
            "valor_total": float(valor_celula(row, headers, "valor_total")),
            "num_parcelas": int(valor_celula(row, headers, "num_parcelas")),
            "valor_parcela": valor_celula(row, headers, "valor_parcela", obrigatorio=False, padrao=None),
            "data_primeira_parcela": pega_data("data_primeira_parcela"),
            "data_emissao": pega_data("data_emissao", padrao=date.today()),
            "local": valor_celula(row, headers, "local", obrigatorio=False, padrao=PADRAO_LOCAL),
            "foro": valor_celula(row, headers, "foro", obrigatorio=False, padrao=PADRAO_FORO),
            "juros_mora_pct": valor_celula(row, headers, "juros_mora_pct", obrigatorio=False, padrao=PADRAO_JUROS_MORA_PCT),
            "multa_moratoria_pct": valor_celula(row, headers, "multa_moratoria_pct", obrigatorio=False, padrao=PADRAO_MULTA_MORATORIA_PCT),
            "multa_pct": valor_celula(row, headers, "multa_pct", obrigatorio=False, padrao=PADRAO_MULTA_PCT),
            "honorarios_pct": valor_celula(row, headers, "honorarios_pct", obrigatorio=False, padrao=PADRAO_HONORARIOS_PCT),
            "desconto_rendimentos_pct": valor_celula(row, headers, "desconto_rendimentos_pct", obrigatorio=False, padrao=PADRAO_DESCONTO_RENDIMENTOS_PCT),
            "testemunha1_nome": valor_celula(row, headers, "testemunha1_nome", obrigatorio=False, padrao=PADRAO_TESTEMUNHA1_NOME),
            "testemunha1_cpf": valor_celula(row, headers, "testemunha1_cpf", obrigatorio=False, padrao=PADRAO_TESTEMUNHA1_CPF),
            "testemunha2_nome": valor_celula(row, headers, "testemunha2_nome", obrigatorio=False, padrao=PADRAO_TESTEMUNHA2_NOME),
            "testemunha2_cpf": valor_celula(row, headers, "testemunha2_cpf", obrigatorio=False, padrao=PADRAO_TESTEMUNHA2_CPF),
        }
        clientes.append(cliente)
    return clientes


# ---------------------------------------------------------------------------
# Calculo das parcelas
# ---------------------------------------------------------------------------

def calcular_parcelas(cliente):
    n = cliente["num_parcelas"]
    total = round(cliente["valor_total"], 2)
    valor_informado = cliente["valor_parcela"]

    if valor_informado:
        valor_base = round(float(valor_informado), 2)
    else:
        valor_base = round(total / n, 2)

    valores = [valor_base] * n
    diferenca = round(total - valor_base * n, 2)
    if diferenca != 0:
        # ajusta a ultima parcela para o total bater certinho
        valores[-1] = round(valores[-1] + diferenca, 2)

    parcelas = []
    for i in range(n):
        vencimento = add_meses(cliente["data_primeira_parcela"], i)
        parcelas.append({
            "numero": i + 1,
            "valor": valores[i],
            "vencimento": vencimento,
        })
    return parcelas


# ---------------------------------------------------------------------------
# Geracao do texto do contrato
# ---------------------------------------------------------------------------

def montar_clausula_parcelas(parcelas, n):
    letras = "abcdefghijklmnopqrstuvwxyz"
    linhas = []
    for p, letra in zip(parcelas, letras):
        final = "." if p["numero"] == n else ";"
        linhas.append(
            f"{letra}) Nota Promissória {p['numero']:02d}/{n:02d} – {formata_moeda(p['valor'])} – "
            f"vencimento em {formata_data(p['vencimento'])}{final}"
        )
    return linhas


def montar_texto_contrato(cliente, parcelas):
    n = cliente["num_parcelas"]
    n_extenso = numero_por_extenso(n)
    nome_devedor = formatar_nome_proprio(cliente["nome_devedor"])
    qualificacao_devedor = (
        f"{nome_devedor}, {cliente['nacionalidade']}, {cliente['estado_civil']}, "
        f"{cliente['profissao']}, portador(a) do CPF nº {cliente['cpf_devedor']}, "
        f"residente e domiciliado(a) na {cliente['endereco_residencial']}"
    )
    if cliente["endereco_profissional"]:
        qualificacao_devedor += f", possuindo domicílio profissional na {cliente['endereco_profissional']}"
    qualificacao_devedor += "."

    clausulas_parcelas = "\n".join(montar_clausula_parcelas(parcelas, n))

    texto = f"""INSTRUMENTO PARTICULAR DE CONFISSÃO DE DÍVIDA, PARCELAMENTO E OUTRAS AVENÇAS

CREDOR: {CREDOR_NOME}, inscrita no CNPJ nº {CREDOR_CNPJ}, com sede na {CREDOR_ENDERECO}, neste ato representada por seu sócio {CREDOR_SOCIO_NOME}, {CREDOR_SOCIO_QUALIFICACAO}, portador do CPF nº {CREDOR_SOCIO_CPF} e RG nº {CREDOR_SOCIO_RG}.

DEVEDOR: {qualificacao_devedor}

As partes acima qualificadas celebram o presente INSTRUMENTO PARTICULAR DE CONFISSÃO DE DÍVIDA, PARCELAMENTO E OUTRAS AVENÇAS, que se regerá pelas cláusulas e condições seguintes.

CLÁUSULA PRIMEIRA – DA ORIGEM, CERTEZA E LIQUIDEZ DA DÍVIDA

1.1. O DEVEDOR reconhece expressamente que é legítimo devedor do CREDOR em razão dos serviços de consultoria e assessoria financeira prestados pelo CREDOR.

1.2. O DEVEDOR confessa, de forma irrevogável e irretratável, ser devedor da quantia líquida, certa, exigível e incontroversa de {formata_moeda(cliente['valor_total'])} ({valor_por_extenso(cliente['valor_total'])}).

1.3. O DEVEDOR declara que firma o presente instrumento de livre e espontânea vontade, sem qualquer vício de consentimento, coação, erro, dolo ou estado de perigo, renunciando expressamente a alegações futuras nesse sentido.

1.4. A presente confissão de dívida constitui novação apenas quanto à forma de pagamento, permanecendo íntegros os direitos creditórios do CREDOR até a quitação integral da obrigação.

CLÁUSULA SEGUNDA – DA FORMA DE PAGAMENTO

2.1. A dívida será liquidada mediante o pagamento de {n:02d} ({n_extenso}) parcelas sucessivas no valor de {formata_moeda(parcelas[0]['valor'])} ({valor_por_extenso(parcelas[0]['valor'])}) cada, representadas por Notas Promissórias emitidas pelo DEVEDOR, com os seguintes vencimentos:

{clausulas_parcelas}

2.2. As Notas Promissórias são emitidas em caráter de garantia da obrigação ora confessada, possuindo natureza pro solvendo, não implicando novação da dívida.

2.3. O pagamento de cada parcela será realizado mediante PIX para a chave CNPJ nº {CREDOR_PIX}, de titularidade do CREDOR, até o final do dia do respectivo vencimento.

2.4. A quitação de cada parcela ocorrerá mediante a efetiva compensação do valor na conta do CREDOR, obrigando-se este a devolver ao DEVEDOR a respectiva Nota Promissória quitada.

CLÁUSULA TERCEIRA – DO VENCIMENTO ANTECIPADO

3.1. O inadimplemento de qualquer parcela, ainda que parcial, acarretará automaticamente o vencimento antecipado das parcelas vincendas, tornando imediatamente exigível o saldo remanescente da dívida.

3.2. Verificada a mora, independentemente de notificação judicial ou extrajudicial, poderá o CREDOR promover a cobrança integral do débito por todos os meios admitidos em direito.

CLÁUSULA QUARTA – DOS ENCARGOS DO INADIMPLEMENTO

4.1. Em caso de atraso no pagamento de qualquer parcela, incidirão sobre o débito:

I – Correção monetária pela Tabela Prática do Tribunal de Justiça do Estado de São Paulo – TJSP, ou índice que venha a substitui-la;

II – Juros moratórios de {percentual_por_extenso(cliente['juros_mora_pct'])} ao mês, calculados pro rata die.

4.2. O inadimplemento de qualquer parcela implicará a resolução antecipada do parcelamento concedido e a incidência de multa penal compensatória correspondente a {percentual_por_extenso(cliente['multa_pct'])} sobre o saldo devedor vencido e antecipadamente vencido, em razão da quebra do acordo firmado entre as partes.

4.3. A multa prevista nesta cláusula possui natureza penal compensatória decorrente da rescisão do parcelamento por culpa exclusiva do DEVEDOR, não se confundindo com os juros moratórios e a atualização monetária.

4.4. Em caso de cobrança judicial, o DEVEDOR responderá ainda pelas custas processuais, despesas de cobrança e honorários advocatícios contratuais equivalentes a {percentual_por_extenso(cliente['honorarios_pct'])} sobre o débito atualizado.

4.5. Sem prejuízo do disposto nos itens anteriores, o atraso no pagamento de qualquer parcela sujeitará o DEVEDOR ao pagamento de multa moratória de {percentual_por_extenso(cliente['multa_moratoria_pct'])} sobre o valor da parcela em atraso, acrescida de juros de mora de {percentual_por_extenso(cliente['juros_mora_pct'])} ao mês e correção monetária pela Tabela Prática do TJSP, ambos calculados pro rata die a partir do dia seguinte ao do vencimento até a data da efetiva liquidação.

CLÁUSULA QUINTA – DA AUTORIZAÇÃO DE DESCONTO EM RENDIMENTOS

5.1. Em caso de mora e vencimento antecipado da obrigação, o DEVEDOR autoriza expressamente que eventual composição amigável ou judicial possa contemplar desconto mensal de até {percentual_por_extenso(cliente['desconto_rendimentos_pct'])} de seus rendimentos líquidos, salários, verbas remuneratórias ou benefícios percebidos, até a integral liquidação da dívida.

5.2. A presente autorização constitui manifestação inequívoca de vontade do DEVEDOR quanto à possibilidade de utilização de parcela de seus rendimentos para satisfação da obrigação, observados os limites legais e eventual determinação judicial quando necessária.

5.3. O DEVEDOR reconhece que a presente obrigação possui origem contratual legítima, líquida e exigível, comprometendo-se a colaborar para sua satisfação integral.

CLÁUSULA SEXTA – DOS ÓRGÃOS DE PROTEÇÃO AO CRÉDITO

6.1. Em caso de inadimplemento, o DEVEDOR autoriza a inscrição e manutenção de seu nome junto aos órgãos de proteção ao crédito, inclusive SPC, SERASA e congêneres, observados os requisitos legais.

CLÁUSULA SÉTIMA – DO TÍTULO EXECUTIVO EXTRAJUDICIAL

7.1. Nos termos do artigo 784, inciso III, do Código de Processo Civil, o presente instrumento constitui título executivo extrajudicial, sendo a dívida líquida, certa e exigível.

7.2. O DEVEDOR reconhece expressamente a exigibilidade do crédito e concorda que o CREDOR promova todas as medidas judiciais cabíveis para satisfação da obrigação em caso de inadimplemento.

CLÁUSULA OITAVA – DA TOLERÂNCIA

8.1. Qualquer tolerância, concessão, atraso ou omissão do CREDOR no exercício de seus direitos não importará novação, renúncia ou alteração contratual, permanecendo íntegros todos os direitos previstos neste instrumento.

CLÁUSULA NONA – DA QUITAÇÃO

9.1. O pagamento integral de todas as parcelas previstas neste instrumento implicará automática, plena, geral, irrevogável e irretratável quitação da dívida ora confessada.

CLÁUSULA DÉCIMA – DO FORO

10.1. As partes elegem o Foro da Comarca de {cliente['foro']} para dirimir quaisquer controvérsias decorrentes deste instrumento, com renúncia expressa a qualquer outro, por mais privilegiado que seja.

E, por estarem justos e contratados, assinam o presente instrumento em conjunto com duas testemunhas.

{cliente['local']}, {formata_data_mes_nome(cliente['data_emissao'])}.


_________________________________________
{CREDOR_NOME_ASSINATURA} – CREDOR


_________________________________________
{nome_devedor} - DEVEDOR


_________________________________________
{cliente['testemunha1_nome']} – CPF {cliente['testemunha1_cpf']} TESTEMUNHA


_________________________________________
{cliente['testemunha2_nome']} – CPF {cliente['testemunha2_cpf']} TESTEMUNHA
"""
    return texto


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

CLAUSULAS_EM_NEGRITO = ("1.3.", "4.2.", "5.1.", "10.1.")


class ContratoPDF(FPDF):
    def header(self):
        pass

    def footer(self):
        pass


def gerar_pdf_contrato(cliente, parcelas, texto_contrato, caminho_saida: Path):
    pdf = ContratoPDF(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(20, 18, 20)
    fonte = registrar_fonte(pdf)
    pdf.add_page()
    pdf.set_font(fonte, size=10.5)
    pdf.set_text_color(0, 0, 0)

    altura_linha = 5.3
    proxima_linha_e_assinatura = False
    primeira_assinatura = True
    # traço + nome de cada signatário ocupam ~17mm; são 4 signatários.
    ALTURA_BLOCO_ASSINATURAS = 72
    for paragrafo in limpar_para_pdf(texto_contrato, fonte).split("\n"):
        if paragrafo.strip() == "":
            pdf.ln(2.5)
            continue

        # Linha de "_____" vira um traço cinza de ponta a ponta (igual ao
        # modelo); a linha seguinte (nome/CPF de quem assina) fica centralizada.
        if set(paragrafo.strip()) == {"_"}:
            # O bloco de assinaturas não se divide entre páginas: se as quatro
            # assinaturas não couberem juntas, começam todas na página seguinte.
            if primeira_assinatura:
                primeira_assinatura = False
                if pdf.h - pdf.b_margin - pdf.get_y() < ALTURA_BLOCO_ASSINATURAS:
                    pdf.add_page()

            # Espaço em branco acima do traço, para caber a assinatura à mão.
            y = pdf.get_y() + 7
            largura_traco = 95
            x_inicio = (pdf.w - largura_traco) / 2
            pdf.set_draw_color(170, 170, 170)
            pdf.set_line_width(0.4)
            pdf.line(x_inicio, y, x_inicio + largura_traco, y)
            pdf.set_draw_color(0, 0, 0)
            pdf.set_line_width(0.2)
            # Cursor logo abaixo do traço (set_y, não ln: o ln avançaria a partir
            # da posição anterior e o nome cairia sobre o traço).
            pdf.set_y(y + 1)
            proxima_linha_e_assinatura = True
            continue

        # No modelo os títulos das cláusulas são em peso normal; o negrito é
        # reservado para as cláusulas que impõem os ônus mais pesados ao devedor.
        negrito = paragrafo.startswith("INSTRUMENTO") or paragrafo.startswith(
            CLAUSULAS_EM_NEGRITO
        )
        pdf.set_font(fonte, style="B" if negrito else "", size=10.5)
        alinhamento = "C" if proxima_linha_e_assinatura else "J"
        proxima_linha_e_assinatura = False

        # Não deixa uma cláusula começar numa página e ser cortada na outra:
        # mede quantas linhas o parágrafo vai ocupar e, se não couber no
        # espaço restante da página atual, pula pra próxima página inteira.
        linhas = pdf.multi_cell(0, altura_linha, paragrafo, align=alinhamento, dry_run=True, output="LINES")
        altura_paragrafo = len(linhas) * altura_linha
        espaco_restante = pdf.h - pdf.b_margin - pdf.get_y()
        if altura_paragrafo > espaco_restante:
            pdf.add_page()

        pdf.multi_cell(0, altura_linha, paragrafo, align=alinhamento)
        pdf.ln(1)

    # --- Notas Promissórias (sempre 3 por página, igual ao modelo original) ---
    n = cliente["num_parcelas"]
    pdf.add_page()
    for idx, p in enumerate(parcelas):
        if idx > 0 and idx % 3 == 0:
            pdf.add_page()
        y0 = pdf.get_y()
        x0 = pdf.l_margin
        largura = pdf.w - pdf.l_margin - pdf.r_margin

        # A caixa cresce quando o endereço do emitente ocupa mais de uma linha,
        # senão o texto encostaria no traço da assinatura.
        texto_endereco = f"Endereço: **{cliente['endereco_residencial'].upper()}**"
        pdf.set_font(fonte, size=9)
        linhas_endereco = pdf.multi_cell(
            largura - 8, 4.6, texto_endereco, markdown=True, dry_run=True, output="LINES"
        )
        altura = 78 + max(0, len(linhas_endereco) - 1) * 4.6

        pdf.rect(x0, y0, largura, altura)

        # Título e número ficam próximos um do outro (a largura da célula do
        # título é medida em vez de fixa, senão sobrava um vão entre os dois).
        pdf.set_xy(x0 + 4, y0 + 4)
        pdf.set_font(fonte, style="B", size=13)
        titulo = "NOTA PROMISSÓRIA"
        largura_titulo = pdf.get_string_width(titulo) + 4
        pdf.cell(largura_titulo, 8, titulo)

        pdf.set_font(fonte, size=10)
        texto_numero = f"Nº {p['numero']:02d}/{n:02d}"
        largura_numero = pdf.get_string_width(texto_numero) + 6
        pdf.cell(largura_numero, 8, texto_numero)

        # Valor e vencimento ficam alinhados à direita, um embaixo do outro:
        # o valor sobe para a primeira linha e o vencimento desce para a segunda.
        largura_resto = largura - 8 - largura_titulo - largura_numero
        x_resto = x0 + 4 + largura_titulo + largura_numero

        pdf.set_font(fonte, style="B", size=12)
        pdf.cell(largura_resto, 8, formata_moeda(p["valor"]), align="R")

        pdf.set_xy(x_resto, y0 + 13)
        pdf.set_font(fonte, size=10)
        pdf.cell(
            largura_resto, 7,
            f"Vencimento: {formata_data_mes_nome(p['vencimento'])}",
            align="R",
        )

        pdf.set_xy(x0 + 4, y0 + 22)
        pdf.set_font(fonte, size=9.5)
        extenso_venc = data_por_extenso(p["vencimento"])
        dia_ext, mes_ext, ano_ext = extenso_venc.split(" de ")
        texto_nota = (
            f"No dia **{dia_ext}** de **{mes_ext}** de **{ano_ext}** pagaremos por esta única via de "
            f"**NOTA PROMISSÓRIA** a **{CREDOR_NOME_NOTA}** CNPJ **{CREDOR_CNPJ}** ou à sua ordem a "
            f"quantia de **{valor_por_extenso(p['valor']).upper()}** em moeda corrente desse país"
        )
        pdf.multi_cell(largura - 8, 4.6, texto_nota, markdown=True)

        pdf.set_xy(x0 + 4, y0 + 44)
        pdf.set_font(fonte, size=9)
        rotulo_local = "Local de pagamento: "
        valor_local = f"{cliente['local'].upper()} - SP"
        pdf.set_font(fonte, style="", size=9)
        largura_local = pdf.get_string_width(rotulo_local)
        pdf.set_font(fonte, style="B", size=9)
        largura_local += pdf.get_string_width(valor_local)
        pdf.set_font(fonte, style="", size=9)
        pdf.cell(largura_local + 6, 5, f"{rotulo_local}**{valor_local}**", markdown=True)
        pdf.cell(
            0, 5,
            f"Data da Emissão: **{formata_data(cliente['data_emissao'])}**",
            markdown=True,
        )
        pdf.set_xy(x0 + 4, y0 + 51)
        pdf.cell(
            0, 5,
            f"Nome do Emitente: **{cliente['nome_devedor'].upper()}**, CPF: **{cliente['cpf_devedor']}**",
            markdown=True,
        )
        pdf.set_xy(x0 + 4, y0 + 58)
        pdf.multi_cell(largura - 8, 4.6, texto_endereco, markdown=True)

        largura_traco_nota = 95
        x_traco_nota = x0 + (largura - largura_traco_nota) / 2
        pdf.line(x_traco_nota, y0 + altura - 6, x_traco_nota + largura_traco_nota, y0 + altura - 6)
        pdf.set_xy(x0 + 4, y0 + altura - 5)
        pdf.set_font(fonte, size=9)
        pdf.cell(largura - 8, 5, "Assinatura do Emitente", align="C")

        pdf.set_y(y0 + altura + 4)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(caminho_saida))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def processar_cliente(cliente):
    parcelas = calcular_parcelas(cliente)
    texto = montar_texto_contrato(cliente, parcelas)
    nome_arquivo = f"{slugify(cliente['nome_devedor'])}_{cliente['data_emissao'].strftime('%Y%m%d')}.pdf"
    caminho_saida = PASTA_SAIDA / nome_arquivo
    gerar_pdf_contrato(cliente, parcelas, texto, caminho_saida)
    print(f"[OK] Linha {cliente['linha_planilha']}: {cliente['nome_devedor']} -> {caminho_saida}")


def main():
    args = sys.argv[1:]
    caminho_planilha = PLANILHA_PADRAO
    linha_alvo = None

    i = 0
    while i < len(args):
        if args[i] == "--planilha":
            caminho_planilha = Path(args[i + 1])
            i += 2
        else:
            linha_alvo = int(args[i])
            i += 1

    if not caminho_planilha.exists():
        print(f"Planilha não encontrada: {caminho_planilha}")
        sys.exit(1)

    clientes = carregar_clientes(caminho_planilha)
    if linha_alvo is not None:
        clientes = [c for c in clientes if c["linha_planilha"] == linha_alvo]
        if not clientes:
            print(f"Nenhuma linha {linha_alvo} encontrada na planilha.")
            sys.exit(1)

    for cliente in clientes:
        try:
            processar_cliente(cliente)
        except Exception as e:
            print(f"[ERRO] Linha {cliente['linha_planilha']} ({cliente.get('nome_devedor')}): {e}")


if __name__ == "__main__":
    main()
